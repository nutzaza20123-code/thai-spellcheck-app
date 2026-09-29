#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
docx_spellcheck.py
====================
เครื่องมือตรวจการสะกดคำ (ไทย + อังกฤษ) แบบละเอียด สำหรับไฟล์ Microsoft Word (.docx)

ออกแบบมาสำหรับตรวจเอกสารสำคัญก่อนส่ง โดยจะ:
  1. อ่านข้อความ "ทุกส่วน" ของไฟล์ .docx ที่เข้าถึงได้ผ่านไลบรารี python-docx ได้แก่
     เนื้อหาหลัก, ตาราง (รวมตารางซ้อนในตาราง), หัวกระดาษ/ท้ายกระดาษทุก section,
     และอ่าน (เฉพาะข้อความ) เชิงอรรถ/อ้างอิงท้ายเรื่องแบบขั้นสูงเพิ่มเติม
  2. ตรวจคำภาษาไทยเทียบกับพจนานุกรมหลายชุด (ฉบับราชบัณฑิตยสถาน + คลังคำทั่วไป +
     ชื่อบุคคล/จังหวัด/ชื่อเฉพาะ) และตรวจคำภาษาอังกฤษด้วย pyspellchecker
  3. ตรวจรูปแบบอื่น ๆ ที่มักพบในเอกสาร: คำซ้ำติดกัน, ตัวอักษรซ้ำผิดปกติ,
     เครื่องหมายวรรคตอนซ้ำ, เว้นวรรคเกิน, ภาษาไทย-อังกฤษ/ตัวเลขติดกันไม่เว้นวรรค
  4. ให้ "ระดับความเชื่อมั่น" 3 ระดับ (สูง/กลาง/ต่ำ) ต่อทุกจุดที่ตรวจพบ พร้อมคำแนะนำ
  5. สร้างผลลัพธ์ 2 ไฟล์:
       - ไฟล์ Word สำเนา ที่ไฮไลต์คำน่าสงสัย พร้อมคอมเมนต์คำแนะนำ (เปิดใน Word ได้ปกติ)
       - ไฟล์ Excel รายงานสรุปทุกจุดที่พบ พร้อมบริบทและคำแนะนำ

ข้อจำกัดที่ควรทราบ (โปรดอ่าน):
  - การตัดคำภาษาไทยไม่มีช่องว่างระหว่างคำ ทำให้การตัดคำมีความคลุมเครือได้เสมอ
    สคริปต์นี้ใช้การตัดคำ 2 อัลกอริทึมไขว้กันเพื่อลดการแจ้งเตือนที่ผิดพลาด
    (false positive) แต่ก็ยังอาจพลาดหรือแจ้งเกินได้บ้าง จึงควรใช้เป็น "ตัวช่วยกรอง"
    ก่อนที่คนจะตรวจทานอีกครั้ง ไม่ใช่เครื่องมือที่แม่นยำ 100%
  - จุดอ่อนที่สำคัญ: ถ้าคำที่สะกดผิดบังเอิญถูกตัดแบ่งออกเป็นคำย่อยที่ "ถูกต้อง" พอดี
    ทั้งคู่ (เช่น พิมพ์ "อนุมัด" ผิด แล้วระบบตัดเป็น "อนุ" + "มัด" ซึ่งทั้งสองคำนี้
    มีอยู่จริงในพจนานุกรม) ระบบจะ "มองไม่เห็น" คำผิดนั้นเลย เพราะเคยทดลองใช้วิธี
    ตรวจคำที่ตัดติดกันเพิ่มเติมแล้วพบว่าทำให้แจ้งเตือนเกินจริงจำนวนมากในข้อความ
    ปกติทั่วไป จึงตัดสินใจไม่ใส่กลไกนั้น เพื่อรักษาความน่าเชื่อถือของรายงาน
    ผลคือ: คำที่ยิ่ง "ผิดแบบมีความหมายบังเอิญ" ยิ่งมีโอกาสหลุดรอดมากกว่าคำที่ผิด
    แบบไม่มีความหมายเลย จึงยังจำเป็นต้องมีคนอ่านทวนเอกสารสำคัญอีกครั้งเสมอ
  - ไม่รองรับข้อความในกล่องข้อความ (text box), SmartArt, หรือวัตถุฝังอื่น ๆ
  - เชิงอรรถ/อ้างอิงท้ายเรื่องจะถูกตรวจและรายงานใน Excel เท่านั้น
    (ไม่สามารถไฮไลต์ย้อนกลับเข้าไฟล์ Word ได้ เนื่องจากข้อจำกัดของไลบรารี)
  - คำที่พบในหัวกระดาษ/ท้ายกระดาษ จะถูก "ไฮไลต์สี" ในไฟล์ Word ให้ตามปกติ แต่จะ
    "ไม่มีคอมเมนต์คำแนะนำ" กำกับ (ทดสอบแล้วพบว่าการใส่คอมเมนต์ในส่วนนี้ทำให้ไฟล์
    เปิดไม่ได้) รายละเอียดคำแนะนำสำหรับจุดเหล่านี้ให้ดูในรายงาน Excel แทน
  - ชื่อเฉพาะ/ศัพท์เทคนิค/ชื่อบริษัทที่ไม่อยู่ในพจนานุกรม จะถูกตีธงเป็น "น่าสงสัย"
    ได้เสมอ แนะนำให้ใส่คำเหล่านี้ในไฟล์พจนานุกรมเสริม (--extra-dict) เพื่อไม่ให้
    ถูกแจ้งซ้ำ ๆ

วิธีใช้งาน:
  python3 docx_spellcheck.py เอกสาร.docx
  python3 docx_spellcheck.py เอกสาร.docx -o ./ผลลัพธ์ --extra-dict ศัพท์เฉพาะ.txt
  python3 docx_spellcheck.py เอกสาร.docx --lang th        # ตรวจเฉพาะภาษาไทย
  python3 docx_spellcheck.py เอกสาร.docx --no-docx         # เอาเฉพาะรายงาน Excel

ต้องติดตั้งไลบรารีก่อน (ดู requirements.txt):
  pip install python-docx pythainlp pyspellchecker openpyxl
"""

from __future__ import annotations

import argparse
import copy
import datetime
import os
import re
import sys
import zipfile
from dataclasses import dataclass, field
from typing import Callable, Iterable, Optional

try:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from docx.text.run import Run
    from docx.text.paragraph import Paragraph
except ImportError:
    sys.exit("ขาดไลบรารี python-docx: กรุณารัน  pip install python-docx")

try:
    from pythainlp.tokenize import word_tokenize
    from pythainlp.corpus import (
        thai_orst_words,
        thai_words,
        thai_male_names,
        thai_female_names,
        thai_family_names,
        provinces,
        thai_wikipedia_titles,
    )
except ImportError:
    sys.exit("ขาดไลบรารี pythainlp: กรุณารัน  pip install pythainlp")

try:
    from spellchecker import SpellChecker
except ImportError:
    sys.exit("ขาดไลบรารี pyspellchecker: กรุณารัน  pip install pyspellchecker")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("ขาดไลบรารี openpyxl: กรุณารัน  pip install openpyxl")


# ============================================================================
# ค่าคงที่ / regex
# ============================================================================

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

# ช่วงยูนิโค้ดตัวอักษรไทย (พยัญชนะ สระ วรรณยุกต์ ฯ ๆ) ไม่รวมเลขไทย
THAI_LETTERS = r"ก-ฺเ-๎"
THAI_DIGITS = r"๐-๙"

CHUNK_RE = re.compile(
    rf"[{THAI_LETTERS}]+"                              # กลุ่มตัวอักษรไทยติดกัน
    rf"|[A-Za-z]+(?:['’\-][A-Za-z]+)*"             # คำอังกฤษ (รองรับ apostrophe/ยัติภังค์ในคำ)
    rf"|[0-9{THAI_DIGITS}]+(?:[.,][0-9{THAI_DIGITS}]+)*"  # ตัวเลข (รองรับทศนิยม/คั่นหลักพัน)
    r"|\s+"                                              # ช่องว่าง/แท็บ/ขึ้นบรรทัด
    r"|.",                                                # อื่น ๆ ทีละตัวอักษร (เครื่องหมายวรรคตอน ฯลฯ)
    re.UNICODE,
)

# โทเคนไทยที่ไม่ควรตรวจสะกดเดี่ยว ๆ (เครื่องหมายซ้ำคำ/ไปยาลน้อย ฯลฯ)
THAI_SKIP_TOKENS = {"ๆ", "ฯ", "ฯลฯ", "ฯพณฯ"}

# เศษคำที่เป็นแค่ "พยัญชนะ + การันต์" (เช่น "ร์", "ต์", "นด์") มักเกิดจากการตัดคำ
# ผิดพลาดตรงคำทับศัพท์/ชื่อเฉพาะที่ไม่รู้จัก (เช่น ซินเนอร์จี้ -> ...เนอ + ร์)
# ไม่ใช่คำที่ยืนคำเดียวได้จริง จึงต้องรวมเข้ากับคำก่อนหน้าก่อนตรวจสะกด
_KARAN_FRAGMENT_RE = re.compile(r"^[ก-ฮ]{1,3}์$")

CONFIDENCE_HIGH = "สูง"
CONFIDENCE_MED = "กลาง"
CONFIDENCE_LOW = "ต่ำ"

# สีไฮไลต์ตามระดับความเชื่อมั่น/ประเภท (ใช้เฉพาะรายการที่ไฮไลต์กลับเข้า docx)
HL_SPELL_HIGH = WD_COLOR_INDEX.YELLOW
HL_SPELL_MED = WD_COLOR_INDEX.TURQUOISE
HL_DUPLICATE = WD_COLOR_INDEX.PINK
HL_REPEATCHAR = WD_COLOR_INDEX.BRIGHT_GREEN

MAX_SUGGESTIONS = 5
MAX_CONTEXT_CHARS = 60  # จำนวนตัวอักษรบริบทซ้าย/ขวาที่แสดงในรายงาน


# ============================================================================
# ส่วนที่ 1: พจนานุกรม
# ============================================================================


@dataclass
class Dictionaries:
    orst: set
    known_thai: set
    known_thai_by_len: dict
    en_checker: "SpellChecker"
    custom_words: set
    custom_words_sorted: list  # เรียงจากยาวไปสั้น ใช้ "ปกป้อง" คำเสริมไม่ให้ถูกตัดคำแยก


def _group_by_length(words: Iterable[str]) -> dict:
    by_len: dict[int, list] = {}
    for w in words:
        by_len.setdefault(len(w), []).append(w)
    return by_len


def load_dictionaries(
    extra_dict_path: Optional[str] = None,
    verbose: bool = True,
    on_progress: Optional["Callable[[str], None]"] = None,
) -> Dictionaries:
    def report(msg: str):
        if verbose:
            print(msg, file=sys.stderr)
        if on_progress:
            on_progress(msg)

    report("กำลังโหลดพจนานุกรมภาษาไทย/อังกฤษ ...")

    orst = set(thai_orst_words())          # พจนานุกรมฉบับราชบัณฑิตยสถาน (มาตรฐาน)
    general = set(thai_words())            # คลังคำทั่วไป (กว้างกว่า รวมคำทับศัพท์/ภาษาปาก)
    names = (
        set(thai_male_names())
        | set(thai_female_names())
        | set(thai_family_names())
    )
    provinces_set = set(provinces())
    wiki_titles = set(thai_wikipedia_titles())  # ชื่อเฉพาะ/คำศัพท์จากวิกิพีเดียไทย

    custom_words = set()
    if extra_dict_path:
        if not os.path.exists(extra_dict_path):
            report(f"⚠ ไม่พบไฟล์พจนานุกรมเสริม: {extra_dict_path} (ข้ามไป)")
        else:
            with open(extra_dict_path, encoding="utf-8-sig") as f:
                for line in f:
                    w = line.strip()
                    if w and not w.startswith("#"):
                        custom_words.add(w)

    known_thai = orst | general | names | provinces_set | wiki_titles | custom_words | THAI_SKIP_TOKENS
    known_thai_by_len = _group_by_length(known_thai)

    # หมายเหตุสำคัญ: เคยลองส่งพจนานุกรมเสริมเข้าไปเป็น custom_dict ของ pythainlp
    # word_tokenize โดยตรง (ทั้ง engine newmm และ longest) แต่พบว่าทำให้กลไกจัดกลุ่ม
    # "อักขระที่ไม่รู้จัก" ของตัวตัดคำเสียไป (กลายเป็นตัดทีละตัวอักษร เช่น "ตรว" ถูก
    # ตัดเป็น "ต","ร","ว" แทนที่จะเป็นก้อนเดียว) ทำให้รายงานคำผิดแตกเป็นเศษคำไม่มี
    # ความหมายและอ่านยากขึ้นมาก จึงเปลี่ยนมาใช้วิธี "ปกป้อง" คำเสริมด้วยการค้นหาคำ
    # เหล่านี้ในข้อความล่วงหน้าก่อนตัดคำแทน (ดูฟังก์ชัน _protect_custom_words)
    custom_words_sorted = sorted((w for w in custom_words if len(w) >= 2), key=len, reverse=True)

    en_checker = SpellChecker()
    if custom_words:
        en_checker.word_frequency.load_words([w.lower() for w in custom_words])

    report(
        f"  พจนานุกรมราชบัณฑิตยสถาน {len(orst):,} คำ | "
        f"คลังคำทั่วไป {len(general):,} คำ | ชื่อบุคคล/จังหวัด {len(names) + len(provinces_set):,} รายการ | "
        f"ชื่อเฉพาะวิกิพีเดีย {len(wiki_titles):,} รายการ | พจนานุกรมเสริม {len(custom_words)} คำ"
    )

    return Dictionaries(
        orst=orst,
        known_thai=known_thai,
        known_thai_by_len=known_thai_by_len,
        en_checker=en_checker,
        custom_words=custom_words,
        custom_words_sorted=custom_words_sorted,
    )


# ============================================================================
# ส่วนที่ 2: ตัวช่วยแนะนำคำที่ถูกต้อง (suggestion)
# ============================================================================


def _candidates_by_length(word: str, by_len: dict, span: int = 2) -> list:
    out = []
    n = len(word)
    for L in range(max(1, n - span), n + span + 1):
        out.extend(by_len.get(L, ()))
    return out


_suggestion_cache: dict[tuple[str, str], list] = {}


def suggest_thai(word: str, dicts: Dictionaries) -> list[str]:
    cache_key = ("th", word)
    if cache_key in _suggestion_cache:
        return _suggestion_cache[cache_key]

    import difflib

    results: list[str] = []

    # 1) ใช้ตัวตรวจของ pythainlp เอง (อิงคลังคำ TNC + ระยะแก้ไขตัวอักษร)
    try:
        from pythainlp.spell import spell as th_spell

        for cand in th_spell(word)[: MAX_SUGGESTIONS]:
            if cand not in results:
                results.append(cand)
    except Exception:
        pass

    # 2) difflib เทียบกับพจนานุกรมของเรา (กรองด้วยความยาวใกล้เคียงก่อนเพื่อความเร็ว)
    if len(results) < MAX_SUGGESTIONS:
        candidates = _candidates_by_length(word, dicts.known_thai_by_len)
        for cand in difflib.get_close_matches(word, candidates, n=MAX_SUGGESTIONS, cutoff=0.6):
            if cand not in results:
                results.append(cand)

    results = results[:MAX_SUGGESTIONS]
    _suggestion_cache[cache_key] = results
    return results


def suggest_english(word: str, dicts: Dictionaries) -> list[str]:
    cache_key = ("en", word.lower())
    if cache_key in _suggestion_cache:
        return _suggestion_cache[cache_key]

    results: list[str] = []
    lw = word.lower()
    best = dicts.en_checker.correction(lw)
    if best and best != lw:
        results.append(best)
    for cand in sorted(dicts.en_checker.candidates(lw) or []):
        if cand != lw and cand not in results:
            results.append(cand)
        if len(results) >= MAX_SUGGESTIONS:
            break

    _suggestion_cache[cache_key] = results
    return results


# ============================================================================
# ส่วนที่ 3: แบบจำลองผลลัพธ์การตรวจ 1 จุด
# ============================================================================


@dataclass
class Finding:
    location: str            # เช่น "เนื้อหา ย่อหน้า 12", "ตาราง 1 แถว 2 คอลัมน์ 3"
    word: str
    category: str             # ประเภทปัญหา
    confidence: str           # สูง / กลาง / ต่ำ
    suggestions: list = field(default_factory=list)
    context: str = ""
    note: str = ""
    start: Optional[int] = None   # ตำแหน่งอักขระเริ่มต้นในย่อหน้า (ใช้ไฮไลต์)
    end: Optional[int] = None
    paragraph: object = None      # อ้างอิงย่อหน้า python-docx (None ถ้าไฮไลต์กลับไม่ได้ เช่น เชิงอรรถ)
    can_highlight: bool = True


def make_context(full_text: str, start: int, end: int, width: int = MAX_CONTEXT_CHARS) -> str:
    left = full_text[max(0, start - width) : start]
    mid = full_text[start:end]
    right = full_text[end : end + width]
    left = ("…" if start - width > 0 else "") + left.replace("\n", " ").replace("\t", " ")
    right = right.replace("\n", " ").replace("\t", " ") + ("…" if end + width < len(full_text) else "")
    return f"{left}[[{mid}]]{right}"


# ============================================================================
# ส่วนที่ 4: ตัดคำ + ตรวจสะกดต่อ 1 บล็อกข้อความ
# ============================================================================


def _chunk(text: str):
    """แบ่งข้อความเป็นกลุ่ม (chunk) พร้อมชนิด: thai/latin/digit/space/other"""
    for m in CHUNK_RE.finditer(text):
        s, e = m.start(), m.end()
        piece = m.group()
        if re.match(rf"^[{THAI_LETTERS}]+$", piece):
            kind = "thai"
        elif re.match(r"^[A-Za-z]", piece):
            kind = "latin"
        elif re.match(rf"^[0-9{THAI_DIGITS}]", piece):
            kind = "digit"
        elif piece.isspace():
            kind = "space"
        else:
            kind = "other"
        yield piece, s, e, kind


def _is_segmentable_into_known_words(sub: str, dicts: Dictionaries) -> bool:
    """เช็คว่าข้อความย่อยนี้ 'ตัดคำใหม่' (newmm) แล้วได้คำที่มีอยู่ในพจนานุกรมล้วน ๆ หรือไม่
    ถ้าใช่ แปลว่าคำที่ธงไว้อาจเป็นผลจากการตัดคำ (segmentation) ไม่ใช่คำสะกดผิดจริง
    -> ลดระดับความเชื่อมั่นลงเป็น 'กลาง' แทน 'สูง'
    """
    if not sub:
        return True
    try:
        toks = [t for t in word_tokenize(sub, engine="newmm") if t.strip()]
    except Exception:
        return False
    if not toks:
        return False
    return all(t in dicts.known_thai for t in toks)


def _find_duplicate_ngrams(
    word_tokens: list[tuple[str, int, int, str]],
    text: str,
    location: str,
    paragraph,
    can_highlight: bool,
    max_n: int = 4,
) -> list[Finding]:
    """หาวลี (1-4 คำ) ที่พิมพ์ซ้ำติดกัน 2 ครั้ง โดยคั่นด้วยช่องว่างเท่านั้น เช่น
    'เรื่องนี้ เรื่องนี้' หรือ 'และ และ' ตรวจแบบวลียาวก่อน เพื่อไม่ให้ซ้ำซ้อนกับ
    การตรวจคำเดี่ยวภายในวลีเดียวกัน"""
    findings: list[Finding] = []
    n = len(word_tokens)
    i = 0
    while i < n:
        matched_len = 0
        upper = min(max_n, (n - i) // 2)
        for gram in range(upper, 0, -1):
            first = word_tokens[i : i + gram]
            second = word_tokens[i + gram : i + 2 * gram]
            if len(second) < gram:
                continue
            gap = text[first[-1][2] : second[0][1]]
            if gap.strip() != "":
                continue
            seq1 = [w if lang == "th" else w.lower() for w, s, e, lang in first]
            seq2 = [w if lang == "th" else w.lower() for w, s, e, lang in second]
            if seq1 == seq2 and any(t.strip() for t in seq1):
                matched_len = gram
                break
        if matched_len:
            first = word_tokens[i : i + matched_len]
            second = word_tokens[i + matched_len : i + 2 * matched_len]
            phrase = text[first[0][1] : first[-1][2]]
            span_start, span_end = first[0][1], second[-1][2]
            findings.append(
                Finding(
                    location=location,
                    word=f"{phrase} {phrase}",
                    category="คำ/วลีซ้ำติดกัน",
                    confidence=CONFIDENCE_HIGH,
                    suggestions=[phrase],
                    context=make_context(text, span_start, span_end),
                    note=f'มีข้อความ "{phrase}" ซ้ำติดกัน 2 ครั้ง น่าจะเกิดจากการพิมพ์ผิด',
                    start=span_start,
                    end=span_end,
                    paragraph=paragraph,
                    can_highlight=can_highlight,
                )
            )
            i += 2 * matched_len
        else:
            i += 1
    return findings


def _protect_custom_words(piece: str, custom_words_sorted: list[str]) -> list[tuple[str, bool]]:
    """แบ่ง piece (ข้อความไทยล้วน) ออกเป็นช่วง ๆ โดยช่วงที่ตรงกับคำในพจนานุกรมเสริม
    ของผู้ใช้แบบเป๊ะ ๆ จะถูกกันไว้ไม่ให้ตัวตัดคำ (word_tokenize) แยกออกจากกัน
    คืนค่าเป็น list ของ (ข้อความย่อย, ถูกป้องกันหรือไม่)"""
    if not custom_words_sorted:
        return [(piece, False)]
    n = len(piece)
    out: list[tuple[str, bool]] = []
    i = 0
    buf_start = 0
    while i < n:
        matched = None
        for w in custom_words_sorted:  # เรียงยาว -> สั้นแล้ว จึงเจอคำยาวสุดที่ match ก่อนเสมอ
            wl = len(w)
            if wl and piece[i : i + wl] == w:
                matched = w
                break
        if matched:
            if i > buf_start:
                out.append((piece[buf_start:i], False))
            out.append((matched, True))
            i += len(matched)
            buf_start = i
        else:
            i += 1
    if buf_start < n:
        out.append((piece[buf_start:n], False))
    return out


def analyze_text_block(
    text: str,
    location: str,
    dicts: Dictionaries,
    paragraph=None,
    check_thai: bool = True,
    check_english: bool = True,
    can_highlight: bool = True,
) -> list[Finding]:
    """ตรวจข้อความ 1 ก้อน (เช่น 1 ย่อหน้า) แล้วคืนรายการ Finding ทั้งหมดที่พบ"""
    findings: list[Finding] = []
    if not text or not text.strip():
        return findings

    chunks = list(_chunk(text))

    # ---------- 4.1 ตรวจคำ (ไทย/อังกฤษ) ----------
    word_tokens: list[tuple[str, int, int, str]] = []  # (word, start, end, lang)

    for piece, s, e, kind in chunks:
        if kind == "thai":
            # ก่อนตัดคำ กันคำในพจนานุกรมเสริมของผู้ใช้ไว้ก่อน ไม่ให้ถูกตัดแยก
            # (ดูเหตุผลที่ไม่ใช้ custom_dict ของ pythainlp ตรง ๆ ในคอมเมนต์ที่ load_dictionaries)
            protected_spans = _protect_custom_words(piece, dicts.custom_words_sorted)
            sub_tokens: list[str] = []
            for sub_piece, is_protected in protected_spans:
                if is_protected:
                    sub_tokens.append(sub_piece)
                    continue
                # ตัดคำด้วย longest ก่อน (จากการทดสอบ ให้ผลลัพธ์ตรงคำจริงมากกว่า newmm ในหลายกรณี)
                try:
                    sub_tokens.extend(word_tokenize(sub_piece, engine="longest"))
                except Exception:
                    sub_tokens.extend(word_tokenize(sub_piece, engine="newmm"))

            # รวมเศษคำ "พยัญชนะ+การันต์" ที่โดดเดี่ยวเข้ากับคำก่อนหน้า (ดูคำอธิบายที่ _KARAN_FRAGMENT_RE)
            merged_tokens: list[str] = []
            for tok in sub_tokens:
                if merged_tokens and _KARAN_FRAGMENT_RE.match(tok):
                    merged_tokens[-1] = merged_tokens[-1] + tok
                else:
                    merged_tokens.append(tok)
            sub_tokens = merged_tokens

            offset = 0
            for tok in sub_tokens:
                tlen = len(tok)
                tstart, tend = s + offset, s + offset + tlen
                offset += tlen
                tok_stripped = tok.strip()
                if tok_stripped and tok_stripped not in THAI_SKIP_TOKENS:
                    word_tokens.append((tok, tstart, tend, "th"))
        elif kind == "latin":
            word_tokens.append((piece, s, e, "en"))

    for word, start, end, lang in word_tokens:
        if lang == "th" and check_thai:
            if len(word) < 1:
                continue
            if word in dicts.known_thai:
                continue
            # ไม่พบในพจนานุกรมใดเลย -> ตรวจว่าตัดคำผิดพลาดหรือไม่
            segmentable = _is_segmentable_into_known_words(word, dicts)
            if segmentable:
                confidence = CONFIDENCE_MED
                note = "ไม่พบคำนี้ในพจนานุกรมโดยตรง แต่ตัดแยกเป็นคำที่ถูกต้องได้ อาจเป็นผลจากการตัดคำ ไม่ใช่คำสะกดผิดเสมอไป โปรดตรวจสอบด้วยตนเอง"
            else:
                confidence = CONFIDENCE_HIGH
                note = "ไม่พบคำนี้ในพจนานุกรมภาษาไทยที่มีทั้งหมด มีแนวโน้มสะกดผิดสูง"
            findings.append(
                Finding(
                    location=location,
                    word=word,
                    category="สะกดผิด (ไทย)",
                    confidence=confidence,
                    suggestions=suggest_thai(word, dicts),
                    context=make_context(text, start, end),
                    note=note,
                    start=start,
                    end=end,
                    paragraph=paragraph,
                    can_highlight=can_highlight,
                )
            )
        elif lang == "en" and check_english:
            if len(word) <= 1:
                continue
            if word.isupper() and len(word) >= 2:
                continue  # ถือว่าเป็นตัวย่อ/อักษรย่อ ไม่ตรวจสะกด
            lw = word.lower()
            if dicts.en_checker.unknown([lw]):
                findings.append(
                    Finding(
                        location=location,
                        word=word,
                        category="สะกดผิด (อังกฤษ)",
                        confidence=CONFIDENCE_HIGH,
                        suggestions=suggest_english(word, dicts),
                        context=make_context(text, start, end),
                        note="ไม่พบคำนี้ในพจนานุกรมภาษาอังกฤษ",
                        start=start,
                        end=end,
                        paragraph=paragraph,
                        can_highlight=can_highlight,
                    )
                )

    # ---------- 4.2 คำ/วลีซ้ำติดกัน (ตรวจถึงระดับวลียาวสูงสุด 4 คำ ไม่ใช่แค่คำเดี่ยว) ----------
    findings.extend(_find_duplicate_ngrams(word_tokens, text, location, paragraph, can_highlight))

    # ---------- 4.3 ตัวอักษรซ้ำผิดปกติ (เช่น "ดีมากกกก") ----------
    # ตรวจจากข้อความดิบโดยตรง (ไม่ใช้ word_tokens) เพราะบางครั้งตัวตัดคำจะตัดกลุ่ม
    # อักษรที่ซ้ำกันออกเป็นหลายโทเคน (เช่น "กกก" ถูกตัดเป็น "กก"+"ก") ทำให้เช็คระดับ
    # คำเดี่ยวไม่เจอ จึงต้องเช็คที่ตัวอักษรดิบซึ่งไม่ถูกกระทบจากการตัดคำ
    for m in re.finditer(rf"([{THAI_LETTERS}A-Za-z])\1{{2,}}", text):
        findings.append(
            Finding(
                location=location,
                word=m.group(),
                category="ตัวอักษรซ้ำผิดปกติ",
                confidence=CONFIDENCE_MED,
                suggestions=[],
                context=make_context(text, m.start(), m.end()),
                note=f'มีตัวอักษร "{m.group(1)}" ซ้ำกัน {len(m.group(0))} ครั้งติดกัน ผิดปกติสำหรับเอกสารทั่วไป',
                start=m.start(),
                end=m.end(),
                paragraph=paragraph,
                can_highlight=can_highlight,
            )
        )

    # ---------- 4.4 เครื่องหมายวรรคตอนซ้ำผิดปกติ (รายงานอย่างเดียว ไม่ไฮไลต์) ----------
    for m in re.finditer(r"\.{2}(?!\.)|\.{4,}", text):
        findings.append(
            Finding(
                location=location,
                word=m.group(),
                category="เครื่องหมายวรรคตอนซ้ำ",
                confidence=CONFIDENCE_LOW,
                context=make_context(text, m.start(), m.end()),
                note="จุด (.) ติดกันผิดปกติ (จุดไข่ปลาที่ถูกต้องคือ 3 จุด)",
                start=m.start(),
                end=m.end(),
                paragraph=paragraph,
                can_highlight=False,
            )
        )
    for m in re.finditer(r"([!?,;:])\1+", text):
        findings.append(
            Finding(
                location=location,
                word=m.group(),
                category="เครื่องหมายวรรคตอนซ้ำ",
                confidence=CONFIDENCE_LOW,
                context=make_context(text, m.start(), m.end()),
                note=f'เครื่องหมาย "{m.group(1)}" ซ้ำกัน {len(m.group())} ครั้ง',
                start=m.start(),
                end=m.end(),
                paragraph=paragraph,
                can_highlight=False,
            )
        )

    # ---------- 4.5 เว้นวรรคเกิน (ช่องว่าง 2 ตัวขึ้นไป) ----------
    for m in re.finditer(r"(?<!\n)  +", text):
        findings.append(
            Finding(
                location=location,
                word="(ช่องว่างซ้อน)",
                category="เว้นวรรค/รูปแบบ",
                confidence=CONFIDENCE_LOW,
                context=make_context(text, m.start(), m.end()),
                note=f"พบช่องว่างติดกัน {len(m.group())} ช่อง",
                start=m.start(),
                end=m.end(),
                paragraph=paragraph,
                can_highlight=False,
            )
        )

    # ---------- 4.6 ภาษาไทยติดกับอังกฤษ/ตัวเลขโดยไม่เว้นวรรค ----------
    for i in range(len(chunks) - 1):
        p1, s1, e1, k1 = chunks[i]
        p2, s2, e2, k2 = chunks[i + 1]
        if e1 != s2:
            continue
        pair = {k1, k2}
        if k1 != k2 and "thai" in pair and ("latin" in pair or "digit" in pair):
            findings.append(
                Finding(
                    location=location,
                    word=f"{p1}|{p2}",
                    category="เว้นวรรค/รูปแบบ",
                    confidence=CONFIDENCE_LOW,
                    context=make_context(text, max(0, s1), min(len(text), e2)),
                    note="ข้อความภาษาไทยติดกับภาษาอังกฤษ/ตัวเลขโดยไม่มีการเว้นวรรค",
                    start=s1,
                    end=e2,
                    paragraph=paragraph,
                    can_highlight=False,
                )
            )

    return findings


# ============================================================================
# ส่วนที่ 5: การไฮไลต์ + คอมเมนต์กลับเข้าไฟล์ Word
# ============================================================================


def _split_run(run: Run, offset: int):
    """แบ่ง run ออกเป็น 2 ส่วนที่ตำแหน่งอักขระ offset (คงรูปแบบ/ฟอนต์เดิมไว้ทั้งสองฝั่ง)"""
    text = run.text
    n = len(text)
    if offset <= 0:
        return None, run
    if offset >= n:
        return run, None
    left_text, right_text = text[:offset], text[offset:]
    new_r = copy.deepcopy(run._r)
    run._r.addnext(new_r)
    run.text = left_text
    new_run = Run(new_r, run._parent)
    new_run.text = right_text
    return run, new_run


def _ensure_boundaries(paragraph: Paragraph, positions: Iterable[int]):
    for pos in sorted(set(p for p in positions if p is not None and p > 0)):
        runs = paragraph.runs
        cum = 0
        for run in runs:
            rlen = len(run.text)
            if cum < pos < cum + rlen:
                _split_run(run, pos - cum)
                break
            cum += rlen


def _runs_between(paragraph: Paragraph, start: int, end: int) -> list[Run]:
    runs = paragraph.runs
    cum = 0
    out = []
    for run in runs:
        rlen = len(run.text)
        rstart, rend = cum, cum + rlen
        if rstart >= start and rend <= end and rlen > 0:
            out.append(run)
        cum = rend
    return out


def _highlight_color_for(finding: Finding):
    if finding.category == "คำ/วลีซ้ำติดกัน":
        return HL_DUPLICATE
    if finding.category == "ตัวอักษรซ้ำผิดปกติ":
        return HL_REPEATCHAR
    if finding.confidence == CONFIDENCE_HIGH:
        return HL_SPELL_HIGH
    return HL_SPELL_MED


def apply_highlights_and_comments(doc: "Document", findings: list[Finding], author: str = "ตัวตรวจอักษรอัตโนมัติ"):
    """ไฮไลต์ + แปะคอมเมนต์คำแนะนำ ลงในย่อหน้าที่เกี่ยวข้อง (กลุ่มตาม paragraph แล้วประมวลผลทีละย่อหน้า)"""
    by_paragraph: dict[int, list[Finding]] = {}
    for f in findings:
        if not f.can_highlight or f.paragraph is None or f.start is None or f.end is None:
            continue
        by_paragraph.setdefault(id(f.paragraph), []).append(f)

    seen_paragraphs = {}
    for f in findings:
        if f.paragraph is not None:
            seen_paragraphs[id(f.paragraph)] = f.paragraph

    applied = 0
    skipped = 0
    comments_skipped_header_footer = 0
    for pid, flist in by_paragraph.items():
        paragraph = seen_paragraphs[pid]
        # หมายเหตุสำคัญ: doc.add_comment() ของ python-docx 1.2.0 ทำให้ไฟล์เสียหาย
        # (เปิดใน Word/LibreOffice ไม่ได้) เมื่อ run ที่อ้างอิงอยู่ในหัวกระดาษ/ท้ายกระดาษ
        # (ยืนยันด้วยการทดสอบจริง) จึงต้องข้ามการคอมเมนต์สำหรับส่วนนี้ แต่ยังคงไฮไลต์ได้ปกติ
        part_type = type(paragraph.part).__name__
        allow_comment = part_type not in ("HeaderPart", "FooterPart")
        # เรียงตามตำแหน่งเริ่มต้น และตัดรายการที่ overlap กันออก (เก็บอันแรกไว้)
        flist_sorted = sorted(flist, key=lambda f: (f.start, f.end))
        chosen = []
        last_end = -1
        for f in flist_sorted:
            if f.start >= last_end:
                chosen.append(f)
                last_end = f.end
            else:
                skipped += 1  # ทับซ้อนกับรายการก่อนหน้า ข้ามการไฮไลต์ (ยังอยู่ใน Excel)

        positions = []
        for f in chosen:
            positions.extend([f.start, f.end])
        try:
            _ensure_boundaries(paragraph, positions)
        except Exception:
            skipped += len(chosen)
            continue

        for f in chosen:
            try:
                runs = _runs_between(paragraph, f.start, f.end)
                if not runs:
                    skipped += 1
                    continue
                color = _highlight_color_for(f)
                for r in runs:
                    r.font.highlight_color = color
                if allow_comment:
                    suggestion_txt = "、".join(f.suggestions) if f.suggestions else "(ไม่มีคำแนะนำอัตโนมัติ)"
                    comment_text = f"[{f.category} | ความเชื่อมั่น: {f.confidence}]\n{f.note}\nคำแนะนำ: {suggestion_txt}"
                    doc.add_comment(runs=runs, text=comment_text, author=author, initials="AI")
                else:
                    comments_skipped_header_footer += 1
                applied += 1
            except Exception:
                skipped += 1
                continue

    return applied, skipped, comments_skipped_header_footer


# ============================================================================
# ส่วนที่ 6: เดินอ่านทั้งเอกสาร (body, ตาราง, หัวกระดาษ/ท้ายกระดาษ, เชิงอรรถ)
# ============================================================================


def iter_table_cells(table, prefix: str):
    for r_idx, row in enumerate(table.rows, 1):
        for c_idx, cell in enumerate(row.cells, 1):
            cell_label = f"{prefix} แถว {r_idx} คอลัมน์ {c_idx}"
            for p_idx, para in enumerate(cell.paragraphs, 1):
                yield f"{cell_label} ย่อหน้า {p_idx}", para
            for nested_idx, nested_table in enumerate(cell.tables, 1):
                yield from iter_table_cells(nested_table, f"{cell_label} > ตารางซ้อน {nested_idx}")


def analyze_document(
    docx_path: str,
    dicts: Dictionaries,
    check_thai: bool = True,
    check_english: bool = True,
    verbose: bool = True,
    on_progress: Optional["Callable[[str], None]"] = None,
) -> tuple["Document", list[Finding]]:
    def report(msg: str):
        if verbose:
            print(msg, file=sys.stderr)
        if on_progress:
            on_progress(msg)

    doc = Document(docx_path)
    findings: list[Finding] = []

    def run_block(label, para):
        text = "".join(r.text for r in para.runs)
        findings.extend(
            analyze_text_block(text, label, dicts, paragraph=para, check_thai=check_thai, check_english=check_english)
        )

    # เนื้อหาหลัก
    report("กำลังตรวจเนื้อหาหลัก ...")
    for i, para in enumerate(doc.paragraphs, 1):
        run_block(f"เนื้อหา ย่อหน้า {i}", para)

    # ตาราง (รวมตารางซ้อน)
    report("กำลังตรวจตาราง ...")
    for t_idx, table in enumerate(doc.tables, 1):
        for label, para in iter_table_cells(table, f"ตาราง {t_idx}"):
            run_block(label, para)

    # หัวกระดาษ / ท้ายกระดาษ ของทุก section
    report("กำลังตรวจหัวกระดาษ/ท้ายกระดาษ ...")
    for s_idx, section in enumerate(doc.sections, 1):
        header_footer_pairs = [
            (section.header, f"หัวกระดาษ (section {s_idx})"),
            (section.footer, f"ท้ายกระดาษ (section {s_idx})"),
            (section.first_page_header, f"หัวกระดาษหน้าแรก (section {s_idx})"),
            (section.first_page_footer, f"ท้ายกระดาษหน้าแรก (section {s_idx})"),
            (section.even_page_header, f"หัวกระดาษหน้าคู่ (section {s_idx})"),
            (section.even_page_footer, f"ท้ายกระดาษหน้าคู่ (section {s_idx})"),
        ]
        for part, label in header_footer_pairs:
            try:
                if part is None or part.is_linked_to_previous:
                    continue
            except Exception:
                pass
            for p_idx, para in enumerate(part.paragraphs, 1):
                run_block(f"{label} ย่อหน้า {p_idx}", para)
            for t_idx, table in enumerate(part.tables, 1):
                for label2, para in iter_table_cells(table, f"{label} ตาราง {t_idx}"):
                    run_block(label2, para)

    # เชิงอรรถ / อ้างอิงท้ายเรื่อง (อ่านข้อความอย่างเดียว ไม่ไฮไลต์กลับ)
    report("กำลังตรวจเชิงอรรถ/อ้างอิงท้ายเรื่อง (ถ้ามี) ...")
    findings.extend(analyze_footnotes_endnotes(docx_path, dicts, check_thai, check_english))

    return doc, findings


def analyze_footnotes_endnotes(docx_path: str, dicts: Dictionaries, check_thai: bool, check_english: bool) -> list[Finding]:
    """อ่านข้อความล้วนจาก word/footnotes.xml และ word/endnotes.xml โดยตรง (python-docx ไม่รองรับส่วนนี้)"""
    from lxml import etree

    findings: list[Finding] = []
    targets = [("word/footnotes.xml", "เชิงอรรถ"), ("word/endnotes.xml", "อ้างอิงท้ายเรื่อง")]

    try:
        with zipfile.ZipFile(docx_path) as z:
            names = set(z.namelist())
            for part_name, label in targets:
                if part_name not in names:
                    continue
                root = etree.fromstring(z.read(part_name))
                note_tag = "footnote" if "footnote" in part_name else "endnote"
                for note in root.findall(f"{{{W_NS}}}{note_tag}"):
                    note_type = note.get(f"{{{W_NS}}}type")
                    if note_type in ("separator", "continuationSeparator"):
                        continue
                    note_id = note.get(f"{{{W_NS}}}id")
                    for p_idx, p in enumerate(note.findall(f"{{{W_NS}}}p"), 1):
                        texts = p.findall(f".//{{{W_NS}}}t")
                        full_text = "".join(t.text or "" for t in texts)
                        if not full_text.strip():
                            continue
                        loc = f"{label} #{note_id} ย่อหน้า {p_idx}"
                        block_findings = analyze_text_block(
                            full_text, loc, dicts, paragraph=None,
                            check_thai=check_thai, check_english=check_english, can_highlight=False,
                        )
                        findings.extend(block_findings)
    except Exception as e:
        print(f"⚠ อ่านเชิงอรรถ/อ้างอิงท้ายเรื่องไม่สำเร็จ: {e}", file=sys.stderr)

    return findings


# ============================================================================
# ส่วนที่ 7: รายงาน Excel
# ============================================================================

CONF_FILL = {
    CONFIDENCE_HIGH: PatternFill(start_color="FFF8CBAD", end_color="FFF8CBAD", fill_type="solid"),
    CONFIDENCE_MED: PatternFill(start_color="FFFFE699", end_color="FFFFE699", fill_type="solid"),
    CONFIDENCE_LOW: PatternFill(start_color="FFD9E1F2", end_color="FFD9E1F2", fill_type="solid"),
}
HEADER_FILL = PatternFill(start_color="FF203864", end_color="FF203864", fill_type="solid")
HEADER_FONT = Font(color="FFFFFFFF", bold=True)
THIN_BORDER = Border(*(Side(style="thin", color="FFB7B7B7"),) * 4)


def build_excel_report(findings: list[Finding], output_path: str, source_filename: str):
    wb = Workbook()

    # -------- ชีตสรุปภาพรวม --------
    ws_summary = wb.active
    ws_summary.title = "สรุปภาพรวม"
    ws_summary["A1"] = "รายงานผลการตรวจอักษร"
    ws_summary["A1"].font = Font(size=16, bold=True)
    ws_summary["A2"] = f"ไฟล์ต้นฉบับ: {source_filename}"
    ws_summary["A3"] = f"วันที่ตรวจ: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws_summary["A4"] = f"จำนวนจุดที่พบทั้งหมด: {len(findings)} จุด"

    row = 6
    ws_summary.cell(row=row, column=1, value="ระดับความเชื่อมั่น").font = Font(bold=True)
    ws_summary.cell(row=row, column=2, value="จำนวน").font = Font(bold=True)
    row += 1
    for level in (CONFIDENCE_HIGH, CONFIDENCE_MED, CONFIDENCE_LOW):
        count = sum(1 for f in findings if f.confidence == level)
        ws_summary.cell(row=row, column=1, value=f"{level} ({'ควรแก้ไข' if level==CONFIDENCE_HIGH else 'ควรตรวจสอบ' if level==CONFIDENCE_MED else 'ข้อสังเกต'})")
        ws_summary.cell(row=row, column=2, value=count).fill = CONF_FILL[level]
        row += 1

    row += 1
    ws_summary.cell(row=row, column=1, value="ประเภทปัญหา").font = Font(bold=True)
    ws_summary.cell(row=row, column=2, value="จำนวน").font = Font(bold=True)
    row += 1
    cat_counts: dict[str, int] = {}
    for f in findings:
        cat_counts[f.category] = cat_counts.get(f.category, 0) + 1
    for cat, cnt in sorted(cat_counts.items(), key=lambda x: -x[1]):
        ws_summary.cell(row=row, column=1, value=cat)
        ws_summary.cell(row=row, column=2, value=cnt)
        row += 1

    ws_summary.column_dimensions["A"].width = 45
    ws_summary.column_dimensions["B"].width = 14

    # -------- ชีตรายละเอียด --------
    ws = wb.create_sheet("รายละเอียดที่พบ")
    headers = ["ลำดับ", "ตำแหน่ง", "คำ/ข้อความที่พบ", "ประเภทปัญหา", "ระดับความเชื่อมั่น", "คำแนะนำ", "บริบท", "หมายเหตุ"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = THIN_BORDER
    ws.freeze_panes = "A2"

    # เรียงลำดับ: ความเชื่อมั่นสูงก่อน แล้วตามตำแหน่งที่พบ
    conf_order = {CONFIDENCE_HIGH: 0, CONFIDENCE_MED: 1, CONFIDENCE_LOW: 2}
    findings_sorted = sorted(findings, key=lambda f: (conf_order.get(f.confidence, 9), f.location))

    for i, f in enumerate(findings_sorted, 1):
        r = i + 1
        values = [
            i,
            f.location,
            f.word,
            f.category,
            f.confidence,
            "、".join(f.suggestions) if f.suggestions else "-",
            f.context,
            f.note,
        ]
        for col, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=col, value=v)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(col in (7, 8)))
            if col == 5:
                cell.fill = CONF_FILL.get(f.confidence, PatternFill())

    widths = {1: 6, 2: 30, 3: 22, 4: 18, 5: 14, 6: 26, 7: 55, 8: 45}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w

    wb.save(output_path)


# ============================================================================
# ส่วนที่ 8: main()
# ============================================================================


def main():
    parser = argparse.ArgumentParser(
        description="ตรวจการสะกดคำไทย/อังกฤษในไฟล์ Word (.docx) แบบละเอียด",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("input", help="ไฟล์ .docx ที่ต้องการตรวจ")
    parser.add_argument("-o", "--output-dir", default=None, help="โฟลเดอร์สำหรับเก็บผลลัพธ์ (ค่าเริ่มต้น: โฟลเดอร์เดียวกับไฟล์ต้นฉบับ)")
    parser.add_argument("--extra-dict", default=None, help="ไฟล์ .txt พจนานุกรมเสริม (คำละ 1 บรรทัด ใช้ # นำหน้าเพื่อคอมเมนต์)")
    parser.add_argument("--lang", choices=["th", "en", "both"], default="both", help="ภาษาที่จะตรวจ (ค่าเริ่มต้น: both)")
    parser.add_argument("--no-docx", action="store_true", help="ไม่ต้องสร้างไฟล์ Word ไฮไลต์ (เอาเฉพาะรายงาน Excel)")
    parser.add_argument("--quiet", action="store_true", help="ไม่ต้องแสดงความคืบหน้า")
    args = parser.parse_args()

    if not os.path.exists(args.input):
        sys.exit(f"ไม่พบไฟล์: {args.input}")
    if not args.input.lower().endswith((".docx",)):
        print(f"⚠ คำเตือน: ไฟล์ '{args.input}' ไม่ใช่นามสกุล .docx จะลองเปิดดู แต่ถ้าเปิดไม่ได้อาจต้องแปลงไฟล์ก่อน", file=sys.stderr)

    verbose = not args.quiet
    check_thai = args.lang in ("th", "both")
    check_english = args.lang in ("en", "both")

    dicts = load_dictionaries(args.extra_dict, verbose=verbose)

    if verbose:
        print(f"กำลังตรวจไฟล์: {args.input}", file=sys.stderr)
    try:
        doc, findings = analyze_document(args.input, dicts, check_thai=check_thai, check_english=check_english, verbose=verbose)
    except Exception as e:
        sys.exit(
            f"ไม่สามารถเปิดไฟล์นี้เป็นเอกสาร Word ได้: {e}\n"
            f"โปรดตรวจสอบว่าเป็นไฟล์ .docx ที่ไม่เสียหาย (ไฟล์ .doc แบบเก่าจะเปิดไม่ได้ "
            f"ต้องแปลงเป็น .docx ก่อน เช่น เปิดด้วย Word แล้วกด 'บันทึกเป็น' เลือก .docx)"
        )

    base = os.path.splitext(os.path.basename(args.input))[0]
    out_dir = args.output_dir or os.path.dirname(os.path.abspath(args.input))
    os.makedirs(out_dir, exist_ok=True)

    excel_path = os.path.join(out_dir, f"{base}_รายงานตรวจคำ.xlsx")
    build_excel_report(findings, excel_path, os.path.basename(args.input))
    if verbose:
        print(f"✓ สร้างรายงาน Excel: {excel_path}", file=sys.stderr)

    docx_out_path = None
    if not args.no_docx:
        applied, skipped, hf_skipped = apply_highlights_and_comments(doc, findings)
        docx_out_path = os.path.join(out_dir, f"{base}_ตรวจแล้ว.docx")
        doc.save(docx_out_path)
        if verbose:
            msg = f"✓ สร้างไฟล์ Word ไฮไลต์: {docx_out_path} (ไฮไลต์แล้ว {applied} จุด, ข้าม {skipped} จุดที่ซ้อนทับ/ไฮไลต์ไม่ได้)"
            if hf_skipped:
                msg += f" [{hf_skipped} จุดในหัวกระดาษ/ท้ายกระดาษ ไฮไลต์ได้แต่ไม่ใส่คอมเมนต์ เนื่องจากข้อจำกัดของไฟล์ Word]"
            print(msg, file=sys.stderr)

    high = sum(1 for f in findings if f.confidence == CONFIDENCE_HIGH)
    med = sum(1 for f in findings if f.confidence == CONFIDENCE_MED)
    low = sum(1 for f in findings if f.confidence == CONFIDENCE_LOW)
    print(f"\nสรุปผล: พบทั้งหมด {len(findings)} จุด (ความเชื่อมั่นสูง {high} / กลาง {med} / ต่ำ {low})")
    if docx_out_path:
        print(f"ไฟล์ Word (ไฮไลต์+คอมเมนต์): {docx_out_path}")
    print(f"ไฟล์รายงาน Excel: {excel_path}")


if __name__ == "__main__":
    main()
